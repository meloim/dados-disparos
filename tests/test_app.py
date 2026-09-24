import hashlib
import hmac
import io
import json
import os
from pathlib import Path
import sqlite3
import tempfile
import time
import unittest
import uuid
from unittest.mock import patch
from openpyxl import load_workbook
from app import create_apps, create_render_app

class Integration(unittest.TestCase):
    def setUp(self):
        self.db = Path(__file__).resolve().parent / ('test-'+uuid.uuid4().hex+'.sqlite3')
        self.settings = {'webhook_secret':'test-only-secret','phone_number_id':'123','aceite':['accept'],'recusa':['reject']}
        panel, webhook = create_apps(self.db,self.settings)
        panel.testing = webhook.testing = True
        self.panel = panel.test_client()
        self.hook = webhook.test_client()
        self.panel.get('/')
        with self.panel.session_transaction() as sess:
            self.csrf = sess['csrf']
    def tearDown(self):
        for suffix in ('','-wal','-shm'):
            Path(str(self.db)+suffix).unlink(missing_ok=True)
    def upload(self, body=None):
        body = body or 'campanha;nome;telefone;mensagem_id\nTeste;Maria;5511999999999;wamid.out\n'
        return self.panel.post('/importar',data={'csrf':self.csrf,'arquivo':(io.BytesIO(body.encode()),'base.csv')})
    def send(self,messages=None,statuses=None,stamp=None,valid=True,channel='123'):
        body = json.dumps({'entry':[{'changes':[{'field':'messages','value':{'metadata':{'phone_number_id':channel},'messages':messages or [],'statuses':statuses or []}}]}]}).encode()
        stamp = str(stamp or int(time.time()))
        signature = 'sha256='+hmac.new(self.settings['webhook_secret'].encode(),stamp.encode()+b'.'+body,hashlib.sha256).hexdigest()
        return self.hook.post('/webhook/datafy',data=body,content_type='application/json',headers={'x-datafy-timestamp':stamp,'x-datafy-signature-256':signature if valid else 'wrong'})
    def msg(self,mid='reply1',context=True,choice='accept',ts=None):
        msg = {'id':mid,'from':'5511999999999','timestamp':str(ts or int(time.time())),'type':'button','button':{'text':'Sim, autorizo','payload':choice}}
        if context: msg['context']={'id':'wamid.out'}
        return msg
    def scalar(self,sql):
        db=sqlite3.connect(self.db)
        try:return db.execute(sql).fetchone()[0]
        finally:db.close()
    def test_signature_replay_and_channel(self):
        self.assertEqual(self.send([self.msg()],valid=False).status_code,401)
        self.assertEqual(self.send([self.msg()],stamp=int(time.time())-700).status_code,401)
        self.assertEqual(self.send([self.msg()],channel='other').status_code,200)
        self.assertEqual(self.scalar('SELECT COUNT(*) FROM events'),0)
        self.assertEqual(self.hook.get('/').status_code,404)
        self.assertEqual(self.panel.post('/importar').status_code,403)
    def test_render_combined_app_requires_login_and_exposes_webhook(self):
        old_password = os.environ.get('PANEL_PASSWORD')
        os.environ['PANEL_PASSWORD'] = 'strong-test-password'
        try:
            cloud = create_render_app(self.db, self.settings)
            cloud.testing = True
            client = cloud.test_client()
            self.assertEqual(client.get('/').status_code,302)
            self.assertEqual(client.get('/healthz').status_code,200)
            # A tela de login precisa do CSS e da logo antes de entrar.
            for asset in ('/static/app.css','/static/logo-construir.svg'):
                with client.get(asset) as response:
                    self.assertEqual(response.status_code,200)
            self.assertEqual(client.post('/webhook/datafy').status_code,401)
            client.get('/login')
            with client.session_transaction() as sess: csrf=sess['csrf']
            self.assertEqual(client.post('/login',data={'csrf':csrf,'username':'admin','password':'wrong'}).status_code,200)
            self.assertEqual(client.post('/login',data={'csrf':csrf,'username':'admin','password':'STRONG-test-password'}).status_code,200)
            self.assertEqual(client.post('/login',data={'csrf':csrf,'username':' Admin ','password':'strong-test-password'}).status_code,302)
            self.assertEqual(client.get('/').status_code,200)
        finally:
            if old_password is None: os.environ.pop('PANEL_PASSWORD',None)
            else: os.environ['PANEL_PASSWORD']=old_password
    def test_idempotence_late_import_and_report(self):
        for _ in range(2):self.assertEqual(self.send([self.msg()]).status_code,200)
        self.assertEqual(self.scalar('SELECT COUNT(*) FROM events'),1)
        self.upload();self.upload()
        self.assertEqual(self.scalar('SELECT COUNT(*) FROM contacts'),1)
        self.assertEqual(self.scalar('SELECT contact_id FROM events'),1)
        for state in ['read','sent','delivered']:
            self.send(statuses=[{'id':'wamid.out','timestamp':str(int(time.time())),'status':state,'recipient_id':'5511999999999'}])
        response=self.panel.get('/relatorio.xlsx')
        self.assertEqual(response.status_code,200)
        book=load_workbook(io.BytesIO(response.data))
        self.assertEqual(book['Contatos']['E2'].value,'Lida')
        self.assertEqual(book['Contatos']['F2'].value,'Aceitou')
        self.assertEqual(book['Resumo']['B5'].value,1)
        self.assertEqual(self.panel.get('/backup').status_code,200)
    def test_unlinked_and_manual_review(self):
        self.upload()
        self.send([self.msg(context=False)])
        self.assertIsNone(self.scalar('SELECT contact_id FROM events'))
        eid=self.scalar('SELECT id FROM events')
        self.assertEqual(self.panel.post('/revisar',data={'csrf':self.csrf,'evento':eid,'contato':'1','resultado':'recusou'}).status_code,302)
        self.assertEqual(self.scalar('SELECT result FROM events'),'recusou')
        self.assertEqual(self.scalar('SELECT COUNT(*) FROM audit'),1)
    def test_text_and_formula_safe_export(self):
        self.upload()
        msg=self.msg();msg['type']='text';msg['text']={'body':'=HYPERLINK("https://example.com")'}
        self.send([msg])
        self.assertEqual(self.scalar('SELECT result FROM events'),'revisar')
        book=load_workbook(io.BytesIO(self.panel.get('/relatorio.xlsx').data))
        self.assertEqual(book['Contatos']['G2'].data_type,'s')
    def test_transactional_import(self):
        self.upload('campanha;nome;telefone;mensagem_id\nTeste;A;5511999999999;wamid.out\nTeste;B;5511888888888;wamid.out\n')
        self.assertEqual(self.scalar('SELECT COUNT(*) FROM contacts'),0)
    def activate(self,name,ts):
        with patch('app.campaign_clock',return_value=ts):
            return self.panel.post('/campanha',data={'csrf':self.csrf,'acao':'ativar','nome':name})
    def status(self,mid,ts,state='sent'):
        return {'id':mid,'timestamp':str(ts),'status':state,'recipient_id':'5511999999999'}
    def test_auto_capture_without_csv_and_duplicates(self):
        now=int(time.time())
        self.activate('Automática',now-10)
        self.send([self.msg()],statuses=[self.status('wamid.out',now)])
        self.send(statuses=[self.status('wamid.out',now),self.status('wamid.second',now)])
        self.assertEqual(self.scalar('SELECT COUNT(*) FROM contacts'),1)
        self.assertEqual(self.scalar('SELECT COUNT(*) FROM outbounds'),2)
        self.assertEqual(self.scalar("SELECT result FROM events WHERE kind='reply'"),'aceitou')
        self.assertEqual(self.scalar("SELECT contact_id FROM events WHERE kind='reply'"),1)
        self.assertEqual(self.panel.get('/').status_code,200)
        book=load_workbook(io.BytesIO(self.panel.get('/relatorio.xlsx').data))
        self.assertEqual(book['Contatos']['F2'].value,'Aceitou')
    def test_late_events_remain_with_original_campaign(self):
        now=int(time.time())
        self.activate('Primeira',now-100)
        with patch('app.campaign_clock',return_value=now-50):
            self.panel.post('/campanha',data={'csrf':self.csrf,'acao':'encerrar'})
        self.activate('Segunda',now-40)
        # Read comes before sent. It cannot identify a campaign by itself.
        self.send(statuses=[self.status('wamid.out',now,'read')])
        self.assertEqual(self.scalar('SELECT COUNT(*) FROM contacts'),0)
        self.send(statuses=[self.status('wamid.out',now-70)])
        self.assertEqual(self.scalar('SELECT campaign FROM contacts'),'Primeira')
        self.assertEqual(self.scalar('SELECT COUNT(*) FROM events WHERE contact_id IS NULL'),0)
        self.send(statuses=[self.status('wamid.new',now)])
        self.assertEqual(self.scalar('SELECT COUNT(*) FROM contacts'),2)
        self.assertEqual(self.scalar('SELECT COUNT(*) FROM campaign_windows WHERE ended IS NULL'),1)
    def test_no_campaign_or_old_send_not_assigned(self):
        now=int(time.time())
        self.send(statuses=[self.status('wamid.old',now-20)])
        self.activate('Nova',now)
        self.send(statuses=[self.status('wamid.old',now,'delivered')])
        self.assertEqual(self.scalar('SELECT COUNT(*) FROM contacts'),0)
        self.activate('Outra',now+1)
        self.assertEqual(self.scalar('SELECT COUNT(*) FROM campaign_windows'),1)
    def test_failed_send_is_captured(self):
        now=int(time.time())
        self.activate('Falhas',now-10)
        self.send(statuses=[self.status('wamid.fail',now,'failed')])
        self.assertEqual(self.scalar('SELECT campaign FROM contacts'),'Falhas')
        self.assertEqual(self.scalar('SELECT COUNT(*) FROM events WHERE contact_id IS NULL'),0)
        self.assertIn('Falhou',self.panel.get('/?campanha=Falhas').get_data(as_text=True))
    def import_list(self,campaign,body):
        return self.panel.post('/importar',data={'csrf':self.csrf,'campanha':campaign,'arquivo':(io.BytesIO(body.encode()),'lista.csv')})
    def status_to(self,mid,phone,ts,state='sent'):
        return {'id':mid,'timestamp':str(ts),'status':state,'recipient_id':phone}
    def test_concurrent_campaigns_by_contact_list(self):
        now=int(time.time())
        # CSV da Datafy como vem: cabeçalho livre, número com nono dígito.
        self.import_list('A','Nome,Número de WhatsApp\nMaria,5581999522801\nJoão,81988343030\n')
        self.import_list('B','telefone;nome\n5571999342429;Ana\n')
        # WhatsApp informa celulares sem o nono dígito.
        self.send(statuses=[self.status_to('wamid.a1','558199522801',now,'failed'),
                            self.status_to('wamid.b1','557199342429',now),
                            self.status_to('wamid.a2','558188343030',now)])
        db=sqlite3.connect(self.db)
        try:campaigns=dict(db.execute('SELECT phone,campaign FROM contacts WHERE outbound IS NOT NULL').fetchall())
        finally:db.close()
        self.assertEqual(campaigns,{'5581999522801':'A','5581988343030':'A','5571999342429':'B'})
        self.assertEqual(self.scalar('SELECT COUNT(*) FROM events WHERE contact_id IS NULL'),0)
        self.assertEqual(self.scalar("SELECT name FROM contacts WHERE phone='5581988343030'"),'João')
    def test_list_imported_after_send_links_old_events(self):
        now=int(time.time())
        self.send(statuses=[self.status_to('wamid.x','558199522801',now,'failed')])
        self.assertEqual(self.scalar('SELECT COUNT(*) FROM contacts'),0)
        self.import_list('Depois','5581999522801\n')
        self.assertEqual(self.scalar('SELECT campaign FROM contacts WHERE outbound IS NOT NULL'),'Depois')
        self.assertIn('Falhou',self.panel.get('/?campanha=Depois').get_data(as_text=True))
    def test_same_phone_goes_to_newest_list(self):
        now=int(time.time())
        self.import_list('Velha','telefone\n5511999999999\n')
        self.send(statuses=[self.status_to('wamid.v','5511999999999',now-100)])
        self.import_list('Nova','telefone\n5511999999999\n')
        self.send(statuses=[self.status_to('wamid.n','5511999999999',now)])
        self.assertEqual(self.scalar("SELECT c.campaign FROM outbounds o JOIN contacts c ON c.id=o.contact_id WHERE o.id='wamid.v'"),'Velha')
        self.assertEqual(self.scalar("SELECT c.campaign FROM outbounds o JOIN contacts c ON c.id=o.contact_id WHERE o.id='wamid.n'"),'Nova')
    def test_move_unlinked_send(self):
        now=int(time.time())
        self.send(statuses=[self.status_to('wamid.m','558199522801',now,'delivered'),self.status_to('wamid.m','558199522801',now,'sent')])
        self.assertIn('wamid.m',self.panel.get('/?aba=config').get_data(as_text=True))
        self.panel.post('/mover',data={'csrf':self.csrf,'campanha':'Manual','envio':['wamid.m']})
        self.assertEqual(self.scalar('SELECT COUNT(*) FROM events WHERE contact_id IS NULL'),0)
        self.assertEqual(self.scalar('SELECT campaign FROM contacts'),'Manual')
    def reply_from(self,mid,phone,ts,text=None,button=None):
        m={'id':mid,'from':phone,'timestamp':str(ts)}
        if button: m.update(type='button',button={'text':button[1],'payload':button[0]})
        else: m.update(type='text',text={'body':text})
        return m
    def test_reply_without_context_links_by_phone_and_is_recognized(self):
        now=int(time.time())
        self.import_list('Convite','nome,telefone\nMaria,5581999522801\n')
        self.send(statuses=[self.status_to('wamid.c','558199522801',now-60)])
        # Escreve direto na conversa (sem "responder"), do WA ID sem o nono dígito.
        self.send([self.reply_from('r1','558199522801',now,'Sim!')])
        self.assertEqual(self.scalar("SELECT result FROM events WHERE kind='reply'"),'aceitou')
        self.assertIsNotNone(self.scalar("SELECT contact_id FROM events WHERE kind='reply'"))
        self.send([self.reply_from('r2','558199522801',now+1,'Não quero, obrigado')])
        self.assertEqual(self.scalar("SELECT result FROM events WHERE id LIKE '%r2'"),'recusou')
    def test_manual_review_accepts_phone_without_ninth_digit(self):
        now=int(time.time())
        self.import_list('Convite','telefone\n5581999522801\n')
        self.send(statuses=[self.status_to('wamid.c','558199522801',now-60)])
        self.send([self.reply_from('r1','558199522801',now,'Pode me explicar melhor como funciona?')])
        eid=self.scalar("SELECT id FROM events WHERE kind='reply'")
        self.assertEqual(self.scalar("SELECT result FROM events WHERE kind='reply'"),'revisar')
        self.assertEqual(self.panel.post('/revisar',data={'csrf':self.csrf,'evento':eid,'resultado':'outra'}).status_code,302)
        self.assertEqual(self.scalar("SELECT result FROM events WHERE kind='reply'"),'outra')
    def test_always_rule_from_button_reclassifies_others(self):
        now=int(time.time())
        self.import_list('Cadastro','telefone\n5511911111111\n5511922222222\n')
        self.send(statuses=[self.status_to('wamid.1','5511911111111',now-60),self.status_to('wamid.2','5511922222222',now-60)])
        self.send([self.reply_from('b1','5511911111111',now,button=('CAD_OK','Participar')),
                   self.reply_from('b2','5511922222222',now,button=('CAD_OK','Participar'))])
        self.assertEqual(self.scalar("SELECT COUNT(*) FROM events WHERE result='revisar'"),2)
        eid=self.scalar("SELECT id FROM events WHERE id LIKE '%b1'")
        self.panel.post('/revisar',data={'csrf':self.csrf,'evento':eid,'resultado':'aceitou','sempre':'1'})
        self.assertEqual(self.scalar("SELECT COUNT(*) FROM events WHERE result='aceitou'"),2)
    def test_failure_reason_and_discard(self):
        now=int(time.time())
        self.import_list('Falhas','telefone\n5581999522801\n')
        failed=self.status_to('wamid.f','558199522801',now,'failed');failed['errors']=[{'code':131026,'title':'Message undeliverable'}]
        self.send(statuses=[failed,self.status_to('wamid.t','5511933333333',now)])
        html=self.panel.get('/?campanha=Falhas').get_data(as_text=True)
        self.assertIn('Número sem WhatsApp',html)
        self.panel.post('/mover',data={'csrf':self.csrf,'acao':'descartar','envio':['wamid.t']})
        self.assertNotIn('wamid.t',self.panel.get('/?aba=config').get_data(as_text=True))
    def test_charts_tab(self):
        now=int(time.time())
        self.assertEqual(self.panel.get('/?aba=graficos').status_code,200)
        self.import_list('A','telefone\n5511911111111\n5511922222222\n')
        self.import_list('B','telefone\n5511933333333\n')
        failed=self.status_to('wamid.3','5511933333333',now-50,'failed');failed['errors']=[{'code':131026}]
        self.send(statuses=[self.status_to('wamid.1','5511911111111',now-600),self.status_to('wamid.1','5511911111111',now-500,'read'),
                            self.status_to('wamid.2','5511922222222',now-600),failed])
        self.send([self.reply_from('r1','5511911111111',now-400,'Sim')])
        html=self.panel.get('/?aba=graficos&periodo=7').get_data(as_text=True)
        self.assertIn('A etapa mais fraca',html)
        self.assertIn('Número sem WhatsApp',html)
        self.assertIn('"answered": 1',html)
        self.assertIn('Taxa de resposta',html)
        self.assertEqual(self.panel.get('/?aba=graficos&periodo=tudo&campanha=B').status_code,200)
    def test_wilson_interval(self):
        from app import wilson
        self.assertEqual(wilson(0,0),[0.0,0.0])
        low,high=wilson(5,10)
        self.assertTrue(low<50<high and 18<low<25 and 75<high<82)
        self.assertEqual(wilson(10,10)[1],100.0)
    def test_reply_implies_read_when_receipts_are_off(self):
        now=int(time.time())
        self.import_list('Sem tique azul','telefone\n5511911111111\n')
        # Só "entregue": a pessoa desligou a confirmação de leitura, mas respondeu.
        self.send(statuses=[self.status_to('wamid.1','5511911111111',now-600),self.status_to('wamid.1','5511911111111',now-590,'delivered')])
        self.send([self.reply_from('r1','5511911111111',now-300,'Sim')])
        html=self.panel.get('/?campanha=Sem+tique+azul').get_data(as_text=True)
        self.assertIn('pela resposta',html)
        charts=self.panel.get('/?aba=graficos&periodo=7').get_data(as_text=True)
        self.assertIn('"label": "Taxa de leitura", "n": 1',charts.replace('"key": "read", ',''))
        self.assertIn('"reads": 0',charts)  # Curva de leitura usa só confirmações reais.
    def test_rules_persist_and_reject_overlap(self):
        self.panel.post('/regras',data={'csrf':self.csrf,'aceite':'YES','recusa':'yes'})
        self.assertEqual(self.scalar('SELECT COUNT(*) FROM preferences'),0)
        self.panel.post('/regras',data={'csrf':self.csrf,'aceite':'custom','recusa':'no'})
        self.send([self.msg(choice='custom')])
        self.assertEqual(self.scalar('SELECT result FROM events'),'aceitou')

    def test_interactive_and_latest_reply(self):
        self.upload()
        msg=self.msg();msg['type']='interactive';msg['interactive']={'type':'button_reply','button_reply':{'id':'reject','title':'Não'}}
        self.send([msg])
        self.send([self.msg(mid='older',ts=int(time.time())-60)])
        book=load_workbook(io.BytesIO(self.panel.get('/relatorio.xlsx').data))
        self.assertEqual(book['Contatos']['F2'].value,'Recusou')

if __name__=='__main__':unittest.main()

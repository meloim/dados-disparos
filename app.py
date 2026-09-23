"""Painel local e receptor isolado de webhooks Datafy. Python 3.11+."""
import csv
import hashlib
import hmac
import io
import json
import os
from pathlib import Path
import re
import secrets
import sqlite3
import threading
import time
import unicodedata
from datetime import datetime, timezone
from contextlib import contextmanager
from zoneinfo import ZoneInfo, ZoneInfoNotFoundError

from flask import Flask, abort, flash, redirect, render_template, request, send_file, session, url_for
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from waitress import serve

try:
    import psycopg
    from psycopg.rows import dict_row
except ImportError:  # Local mode does not require PostgreSQL.
    psycopg = None

ROOT = Path(__file__).resolve().parent
LABELS = {'pendente': 'Sem resposta', 'aceitou': 'Aceitou', 'recusou': 'Recusou', 'revisar': 'Revisar'}

def campaign_clock():
    return int(time.time())

def normalized(value):
    return unicodedata.normalize('NFC', str(value)).strip().casefold()

def phone(value):
    result = re.sub(r'\D', '', str(value))
    if not 10 <= len(result) <= 15:
        raise ValueError('Telefone inválido. Use DDI + DDD + número, como 5511999999999.')
    return result

def date_text(value):
    if not value:
        return ''
    try:
        tz = ZoneInfo('America/Sao_Paulo')
    except ZoneInfoNotFoundError:
        from datetime import timedelta
        tz = timezone(timedelta(hours=-3))
    return datetime.fromtimestamp(int(value), tz).strftime('%d/%m/%Y %H:%M:%S')

def create_apps(db_path=None, settings=None):
    if settings is None:
        cfg = ROOT / 'config.json'
        settings = json.loads(cfg.read_text(encoding='utf-8-sig')) if cfg.exists() else {}
    settings = dict(settings)
    settings['webhook_secret'] = os.environ.get('DATAFY_WEBHOOK_SECRET', settings.get('webhook_secret', '')).strip().strip('"\'')
    settings['phone_number_id'] = str(os.environ.get('DATAFY_PHONE_NUMBER_ID', settings.get('phone_number_id', ''))).strip()
    database_url = os.environ.get('DATABASE_URL', '') if db_path is None else ''
    if database_url.startswith('postgres://'):
        database_url = 'postgresql://' + database_url[len('postgres://'):]
    db_path = Path(db_path or ROOT / 'dados' / 'relatorios.sqlite3')
    if not database_url:
        db_path.parent.mkdir(parents=True, exist_ok=True)

    class Database:
        def __init__(self, raw, postgres=False):
            self.raw, self.postgres = raw, postgres
        def execute(self, sql, params=()):
            if self.postgres:
                if sql.strip().upper() == 'BEGIN IMMEDIATE':
                    sql = 'BEGIN'
                sql = sql.replace('?', '%s')
            return self.raw.execute(sql, params)
        def executescript(self, script):
            if self.postgres:
                for statement in script.split(';'):
                    if statement.strip() and not statement.lstrip().upper().startswith('PRAGMA'):
                        self.raw.execute(statement)
            else:
                self.raw.executescript(script)

    @contextmanager
    def connect():
        if database_url:
            if psycopg is None:
                raise RuntimeError('Instale psycopg para usar PostgreSQL.')
            raw = psycopg.connect(database_url, row_factory=dict_row)
            db = Database(raw, True)
        else:
            raw = sqlite3.connect(db_path, timeout=20)
            raw.row_factory = sqlite3.Row
            raw.execute('PRAGMA foreign_keys=ON')
            db = Database(raw)
        try:
            with raw:
                yield db
        finally:
            raw.close()

    with connect() as db:
        schema = '''
        PRAGMA journal_mode=WAL;
        CREATE TABLE IF NOT EXISTS contacts (
          id INTEGER PRIMARY KEY, campaign TEXT NOT NULL, name TEXT NOT NULL,
          phone TEXT NOT NULL, outbound TEXT UNIQUE,
          UNIQUE(campaign,phone));
        CREATE TABLE IF NOT EXISTS events (
          id TEXT PRIMARY KEY, kind TEXT NOT NULL, phone TEXT, channel TEXT,
          ts INTEGER NOT NULL, context TEXT, body TEXT, choice TEXT, result TEXT,
          contact_id INTEGER REFERENCES contacts(id), association TEXT, raw TEXT NOT NULL);
        CREATE INDEX IF NOT EXISTS events_contact ON events(contact_id,ts);
        CREATE TABLE IF NOT EXISTS audit (
          id INTEGER PRIMARY KEY, ts INTEGER NOT NULL, event_id TEXT, action TEXT NOT NULL);
        CREATE TABLE IF NOT EXISTS campaign_windows (
          id INTEGER PRIMARY KEY, campaign TEXT NOT NULL, started INTEGER NOT NULL, ended INTEGER);
        CREATE UNIQUE INDEX IF NOT EXISTS one_active_window ON campaign_windows((1)) WHERE ended IS NULL;
        CREATE TABLE IF NOT EXISTS outbounds (
          id TEXT PRIMARY KEY, contact_id INTEGER NOT NULL REFERENCES contacts(id));
        INSERT INTO outbounds SELECT outbound,id FROM contacts WHERE outbound IS NOT NULL ON CONFLICT DO NOTHING;
        CREATE TABLE IF NOT EXISTS preferences (key TEXT PRIMARY KEY, value TEXT NOT NULL);
        '''
        if database_url:
            schema = schema.replace('id INTEGER PRIMARY KEY', 'id BIGSERIAL PRIMARY KEY')
        db.executescript(schema)

    panel = Flask(__name__)
    panel.secret_key = os.environ.get('SECRET_KEY') or secrets.token_hex(32)
    panel.config.update(MAX_CONTENT_LENGTH=5*1024*1024, SESSION_COOKIE_HTTPONLY=True,
        SESSION_COOKIE_SAMESITE='Strict', SESSION_COOKIE_SECURE=bool(os.environ.get('RENDER')))
    webhook = Flask('datafy_webhook')
    webhook.config['MAX_CONTENT_LENGTH'] = 2*1024*1024

    @panel.before_request
    def csrf_check():
        public = request.endpoint in ('login', 'health') or request.path == '/webhook/datafy'
        if not public and os.environ.get('PANEL_PASSWORD') and not session.get('authenticated'):
            return redirect(url_for('login', next=request.path))
        if request.method == 'POST' and not secrets.compare_digest(session.get('csrf', ''), request.form.get('csrf', '!')):
            if request.path != '/webhook/datafy':
                abort(403)
        session.setdefault('csrf', secrets.token_hex(24))

    @panel.after_request
    def security_headers(response):
        response.headers['X-Content-Type-Options'] = 'nosniff'
        response.headers['X-Frame-Options'] = 'DENY'
        response.headers['Referrer-Policy'] = 'no-referrer'
        response.headers['Content-Security-Policy'] = "default-src 'self'; style-src 'self' 'unsafe-inline'; form-action 'self'; frame-ancestors 'none'"
        if os.environ.get('RENDER'):
            response.headers['Strict-Transport-Security'] = 'max-age=31536000; includeSubDomains'
        return response

    @panel.route('/login', methods=['GET','POST'])
    def login():
        if request.method == 'POST':
            username = os.environ.get('PANEL_USERNAME', 'admin')
            password = os.environ.get('PANEL_PASSWORD', '')
            ok = password and secrets.compare_digest(request.form.get('username',''), username)
            ok = ok and secrets.compare_digest(request.form.get('password',''), password)
            if ok:
                session.clear()
                session['authenticated'] = True
                session['csrf'] = secrets.token_hex(24)
                return redirect(url_for('index'))
            flash('Usuário ou senha inválidos.')
        return render_template('login.html')

    @panel.post('/sair')
    def logout():
        session.clear()
        return redirect(url_for('login'))

    @panel.get('/healthz')
    def health():
        return {'ok': True}, 200

    def classify(choice):
        val = normalized(choice)
        rules = get_rules()
        yes = {normalized(v) for v in rules.get('aceite', [])}
        no = {normalized(v) for v in rules.get('recusa', [])}
        if val and val in yes and val not in no:
            return 'aceitou'
        if val and val in no and val not in yes:
            return 'recusou'
        return 'revisar'

    def get_rules():
        rules = {k: settings.get(k, []) for k in ('aceite','recusa')}
        with connect() as db:
            for row in db.execute('SELECT * FROM preferences'):
                if row['key'] in rules:
                    rules[row['key']] = json.loads(row['value'])
        return rules

    def link_by_context(db):
        db.execute('''INSERT INTO outbounds SELECT outbound,id FROM contacts WHERE outbound IS NOT NULL ON CONFLICT DO NOTHING''')
        db.execute('''UPDATE events SET contact_id=(SELECT contact_id FROM outbounds WHERE id=events.context),
          association='ID da mensagem enviada' WHERE contact_id IS NULL AND context IS NOT NULL
          AND EXISTS(SELECT 1 FROM outbounds o JOIN contacts ON contacts.id=o.contact_id WHERE o.id=events.context
          AND (events.kind='status' OR contacts.phone=events.phone))''')

    def capture_sends(db):
        # Only the original sent timestamp identifies the campaign window.
        # Delivery/read can happen days later and must never select a new campaign.
        for ev in db.execute("SELECT * FROM events WHERE kind='status' AND body='sent' AND contact_id IS NULL").fetchall():
            if not ev['context'] or not 10 <= len(ev['phone'] or '') <= 15:
                continue
            if db.execute('SELECT 1 FROM outbounds WHERE id=?',(ev['context'],)).fetchone():
                continue
            window = db.execute('''SELECT campaign FROM campaign_windows WHERE started<=?
                AND (ended IS NULL OR ?<ended) ORDER BY started DESC LIMIT 1''',(ev['ts'],ev['ts'])).fetchone()
            if not window:
                continue
            db.execute('''INSERT INTO contacts(campaign,name,phone)
              VALUES(?,?,?) ON CONFLICT(campaign,phone) DO NOTHING''',(window['campaign'],ev['phone'],ev['phone']))
            contact = db.execute('SELECT id FROM contacts WHERE campaign=? AND phone=?',(window['campaign'],ev['phone'])).fetchone()
            db.execute('INSERT INTO outbounds VALUES(?,?) ON CONFLICT DO NOTHING',(ev['context'],contact['id']))
            db.execute('UPDATE contacts SET outbound=COALESCE(outbound,?) WHERE id=?',(ev['context'],contact['id']))

    @webhook.post('/webhook/datafy')
    def receive():
        secret = settings.get('webhook_secret', '')
        channel = str(settings.get('phone_number_id', '')).strip()
        if not secret or not channel:
            return {'error': 'Configure secret e phone_number_id no servidor'}, 503
        raw = request.get_data()
        stamp = request.headers.get('x-datafy-timestamp', '')
        sig = request.headers.get('x-datafy-signature-256', '')
        try:
            if abs(time.time() - int(stamp)) > 300:
                print(f'[webhook 401] timestamp fora da janela: {stamp!r} (agora {int(time.time())})', flush=True)
                abort(401)
        except ValueError:
            print(f'[webhook 401] timestamp ausente/inválido: {stamp!r}', flush=True)
            abort(401)
        expected = 'sha256=' + hmac.new(secret.encode(), stamp.encode()+b'.'+raw, hashlib.sha256).hexdigest()
        if not hmac.compare_digest(expected.encode(), sig.encode()):
            print(f'[webhook 401] assinatura não confere: recebida={sig[:16]!r}... '
                  f'secret_prefixo={secret[:9]!r} secret_len={len(secret)} corpo={len(raw)} bytes', flush=True)
            abort(401)
        try:
            payload = json.loads(raw)
            records = []
            for entry in payload.get('entry', []):
                for change in entry.get('changes', []):
                    value = change.get('value', {})
                    if change.get('field') != 'messages':
                        continue
                    number = str(value.get('metadata', {}).get('phone_number_id', ''))
                    if number != channel:
                        continue
                    for msg in value.get('messages', []):
                        kind = msg.get('type', '')
                        choice = ''
                        body = ''
                        if kind == 'text':
                            body = msg.get('text', {}).get('body', '')
                        elif kind == 'button':
                            button = msg.get('button', {})
                            body = button.get('text', '')
                            choice = button.get('payload') or body
                        elif kind == 'interactive':
                            inter = msg.get('interactive', {})
                            button = inter.get(inter.get('type', ''), {})
                            body = button.get('title', '')
                            choice = button.get('id') or body
                        else:
                            body = '[Mensagem de tipo: '+kind+']'
                        result = classify(choice) if choice else 'revisar'
                        context = msg.get('context', {}).get('id')
                        records.append((number+':message:'+msg['id'], 'reply', phone(msg['from']), number,
                            int(msg['timestamp']), context, str(body), str(choice), result, json.dumps(msg, ensure_ascii=False)))
                    for status in value.get('statuses', []):
                        state = status.get('status', '')
                        if state not in ('sent', 'delivered', 'read', 'failed'):
                            continue
                        records.append((number+':status:'+status['id']+':'+state, 'status',
                            re.sub(r'\D', '', str(status.get('recipient_id', ''))), number,
                            int(status['timestamp']), status['id'], state, '', '', json.dumps(status, ensure_ascii=False)))
            with connect() as db:
                for rec in records:
                    db.execute('''INSERT INTO events
                      (id,kind,phone,channel,ts,context,body,choice,result,raw) VALUES (?,?,?,?,?,?,?,?,?,?)
                      ON CONFLICT(id) DO NOTHING''', rec)
                capture_sends(db)
                link_by_context(db)
            return {'ok': True}, 200
        except (ValueError, KeyError, TypeError, AttributeError):
            return {'error': 'Formato de evento inválido'}, 400

    def report(campaign=''):
        with connect() as db:
            contacts = db.execute("SELECT * FROM contacts WHERE (?='' OR campaign=?) ORDER BY campaign,name,phone", (campaign,campaign)).fetchall()
            events = db.execute('SELECT * FROM events WHERE contact_id IS NOT NULL ORDER BY ts,id').fetchall()
        by_contact = {}
        for ev in events:
            by_contact.setdefault(ev['contact_id'], []).append(dict(ev))
        rows = []
        for item in contacts:
            row = dict(item)
            evs = by_contact.get(row['id'], [])
            replies = [e for e in evs if e['kind']=='reply']
            statuses = {e['body'] for e in evs if e['kind']=='status'}
            row['delivery'] = next((label for code,label in [('read','Lida'),('delivered','Entregue'),('failed','Falhou'),('sent','Enviada')] if code in statuses), 'Não informado')
            row['result'] = replies[-1]['result'] if replies else 'pendente'
            row['answer'] = replies[-1]['body'] if replies else ''
            row['choice'] = replies[-1]['choice'] if replies else ''
            row['when'] = date_text(replies[-1]['ts']) if replies else ''
            row['association'] = replies[-1]['association'] if replies else ''
            rows.append(row)
        totals = {k: sum(r['result']==k for r in rows) for k in LABELS}
        totals.update(total=len(rows), delivered=sum(r['delivery'] in ('Entregue','Lida') for r in rows), read=sum(r['delivery']=='Lida' for r in rows))
        return rows, totals

    @panel.get('/')
    def index():
        campaign = request.args.get('campanha','')
        rows, totals = report(campaign)
        with connect() as db:
            campaigns = [r['campaign'] for r in db.execute('SELECT campaign FROM contacts UNION SELECT campaign FROM campaign_windows ORDER BY campaign')]
            active = db.execute('SELECT * FROM campaign_windows WHERE ended IS NULL').fetchone()
            unlinked = db.execute("SELECT COUNT(DISTINCT context) AS total FROM events WHERE kind='status' AND contact_id IS NULL").fetchone()['total']
            inbox = db.execute("SELECT * FROM events WHERE kind='reply' AND (contact_id IS NULL OR result='revisar') ORDER BY ts DESC LIMIT 200").fetchall()
            last = db.execute('SELECT MAX(ts) AS last_ts FROM events').fetchone()['last_ts']
        return render_template('index.html', rows=rows, totals=totals, campaigns=campaigns,
            campaign=campaign, inbox=inbox, labels=LABELS, date_text=date_text,
            configured=bool(settings.get('webhook_secret') and settings.get('phone_number_id')),
            last=date_text(last), all_contacts=report()[0], rules=get_rules(), active=active,
            unlinked=unlinked, hosted=bool(os.environ.get('PANEL_PASSWORD')))

    @panel.post('/campanha')
    def campaign_control():
        action = request.form.get('acao')
        name = request.form.get('nome','').strip()
        if action not in ('ativar','encerrar') or (action=='ativar' and not 1 <= len(name) <= 150):
            abort(400)
        with connect() as db:
            db.execute('BEGIN IMMEDIATE')
            active = db.execute('SELECT * FROM campaign_windows WHERE ended IS NULL').fetchone()
            if action=='ativar' and active:
                flash('Encerre a campanha ativa antes de iniciar outra.')
            else:
                now = campaign_clock()
                if active:
                    db.execute('UPDATE campaign_windows SET ended=? WHERE id=?',(now,active['id']))
                if action=='ativar':
                    db.execute('INSERT INTO campaign_windows(campaign,started) VALUES(?,?)',(name,now))
                db.execute('INSERT INTO audit(ts,action) VALUES(?,?)',(now,json.dumps({'campaign':name,'action':action})))
                flash('Campanha ativada. Inicie os disparos na Datafy.' if action=='ativar' else 'Captação de novos envios encerrada. Atualizações dos envios registrados continuam chegando.')
        return redirect(url_for('index'))

    @panel.post('/regras')
    def save_rules():
        rules = {k:[v.strip() for v in request.form.get(k,'').splitlines() if v.strip()] for k in ('aceite','recusa')}
        if {normalized(v) for v in rules['aceite']} & {normalized(v) for v in rules['recusa']}:
            flash('Uma opção não pode representar aceite e recusa ao mesmo tempo.')
            return redirect(url_for('index'))
        with connect() as db:
            for k,v in rules.items():
                db.execute('''INSERT INTO preferences(key,value) VALUES(?,?)
                  ON CONFLICT(key) DO UPDATE SET value=excluded.value''',(k,json.dumps(v)))
        flash('Regras salvas para as próximas respostas. Respostas anteriores permanecem disponíveis para revisão.')
        return redirect(url_for('index'))

    @panel.post('/importar')
    def import_csv():
        try:
            upload = request.files.get('arquivo')
            if not upload:
                raise ValueError('Selecione um CSV.')
            raw = upload.read().decode('utf-8-sig')
            dialect = csv.Sniffer().sniff(raw[:4096], delimiters=';,\t')
            reader = csv.DictReader(io.StringIO(raw), dialect=dialect)
            if not {'campanha','nome','telefone'}.issubset(reader.fieldnames or []):
                raise ValueError('Colunas obrigatórias: campanha, nome, telefone. Opcional: mensagem_id.')
            parsed = []
            for n,row in enumerate(reader,2):
                campaign, name = row['campanha'].strip(), row['nome'].strip()
                number = phone(row['telefone'])
                outbound = (row.get('mensagem_id') or '').strip() or None
                if not campaign or not name or len(campaign)>150 or len(name)>200:
                    raise ValueError(f'Linha {n}: nome e campanha precisam estar preenchidos e ter até 200/150 caracteres.')
                if outbound and not outbound.startswith('wamid.'):
                    raise ValueError(f'Linha {n}: mensagem_id deve ser o wamid retornado no envio; deixe vazio se não tiver.')
                parsed.append((campaign,name,number,outbound))
            if not parsed:
                raise ValueError('O arquivo não contém contatos.')
            with connect() as db:
                for campaign,name,number,outbound in parsed:
                    old = db.execute('SELECT outbound FROM contacts WHERE campaign=? AND phone=?',(campaign,number)).fetchone()
                    if old and old['outbound'] and outbound and old['outbound']!=outbound:
                        raise ValueError('Este contato já tem outro envio nesta campanha. Crie uma campanha diferente para novo disparo.')
                    db.execute('''INSERT INTO contacts(campaign,name,phone,outbound) VALUES(?,?,?,?)
                      ON CONFLICT(campaign,phone) DO UPDATE SET name=excluded.name,
                      outbound=COALESCE(contacts.outbound,excluded.outbound)''',(campaign,name,number,outbound))
                link_by_context(db)
            flash(f'{len(parsed)} linhas importadas. Reimportações não duplicam contatos.')
        except (ValueError, UnicodeError, csv.Error, sqlite3.IntegrityError,
                psycopg.IntegrityError if psycopg else sqlite3.IntegrityError) as exc:
            flash('Importação cancelada: '+str(exc))
        return redirect(url_for('index'))

    @panel.post('/revisar')
    def review():
        event_id = request.form.get('evento','')
        result = request.form.get('resultado','')
        if result not in ('aceitou','recusou','revisar'):
            abort(400)
        with connect() as db:
            ev = db.execute("SELECT * FROM events WHERE id=? AND kind='reply'",(event_id,)).fetchone()
            contact = db.execute('SELECT * FROM contacts WHERE id=?',(request.form.get('contato'),)).fetchone()
            if not ev or not contact or ev['phone'] != contact['phone']:
                abort(400)
            db.execute('UPDATE events SET contact_id=?,result=?,association=? WHERE id=?', (contact['id'], result, 'Revisão manual local', event_id))
            db.execute('INSERT INTO audit(ts,event_id,action) VALUES(?,?,?)',(int(time.time()),event_id,json.dumps({'before_contact':ev['contact_id'],'before_result':ev['result'],'contact':contact['id'],'result':result})))
        flash('Revisão salva. A resposta original foi preservada.')
        return redirect(url_for('index'))

    @panel.get('/modelo.csv')
    def model():
        return send_file(io.BytesIO('campanha;nome;telefone\n'.encode('utf-8-sig')), as_attachment=True, download_name='modelo-contatos.csv', mimetype='text/csv')

    @panel.get('/relatorio.xlsx')
    def export():
        def append(ws, values):
            ws.append([re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f]', '', v) if isinstance(v,str) else v for v in values])
        campaign = request.args.get('campanha','')
        rows, totals = report(campaign)
        wb = Workbook()
        summary = wb.active
        summary.title = 'Resumo'
        for row in [['Relatório de autorização — WhatsApp'],['Campanha',campaign or 'Todas'],['Gerado em', date_text(time.time())],['Indicador','Quantidade','Percentual da base'],['Contatos na base',totals['total'],1 if totals['total'] else 0],['Entregues (inclui lidas)',totals['delivered'],totals['delivered']/totals['total'] if totals['total'] else 0],['Leituras informadas',totals['read'],totals['read']/totals['total'] if totals['total'] else 0]]:
            summary.append(row)
        for k,label in LABELS.items():
            summary.append([label,totals[k],totals[k]/totals['total'] if totals['total'] else 0])
        summary.append(['Base importada não comprova envio. Sem status significa não informado.'])
        summary.append(['Respostas sem vínculo não entram nos resultados por campanha.'])
        summary.append(['Resultado considera a resposta vinculada mais recente, inclusive revisão.'])
        for r in range(5,12):
            summary.cell(r,3).number_format = '0.0%'
        detail = wb.create_sheet('Contatos')
        detail.append(['Campanha','Nome','Telefone','ID do envio','Entrega','Resultado','Resposta original','Opção (ID)','Data da resposta','Vínculo'])
        for row in rows:
            append(detail,[row['campaign'],row['name'],row['phone'],row['outbound'],row['delivery'],LABELS[row['result']],row['answer'],row['choice'],row['when'],row['association']])
        history = wb.create_sheet('Histórico')
        history.append(['Campanha','Telefone','Tipo','Data','Conteúdo','Opção','Resultado','Vínculo','ID do evento'])
        with connect() as db:
            evs = db.execute('''SELECT e.*,c.campaign FROM events e LEFT JOIN contacts c ON c.id=e.contact_id
              WHERE (?='' OR c.campaign=?) ORDER BY e.ts''',(campaign,campaign)).fetchall()
        for ev in evs:
            append(history,[ev['campaign'] or 'Sem vínculo',ev['phone'],ev['kind'],date_text(ev['ts']),ev['body'],ev['choice'],LABELS.get(ev['result'],''),ev['association'],ev['id']])
        for ws in wb:
            ws.freeze_panes = 'A2'
            if ws != summary:
                ws.auto_filter.ref = ws.dimensions
            for cell in ws[1]:
                cell.fill = PatternFill('solid', fgColor='164E63')
                cell.font = Font(color='FFFFFF',bold=True)
            for col in ws.columns:
                ws.column_dimensions[col[0].column_letter].width = min(55,max(16,max(len(str(c.value or '')) for c in col)+2))
            # Untrusted text must never become an Excel formula or illegal XML.
            for row in ws:
                for cell in row:
                    if isinstance(cell.value,str):
                        cell.value = re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f]', '', cell.value)
                        cell.data_type = 's'
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return send_file(output, as_attachment=True, download_name='relatorio-datafy.xlsx', mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

    @panel.get('/backup')
    def backup():
        if database_url:
            tables = ('contacts','events','audit','campaign_windows','outbounds','preferences')
            snapshot = {'generated_at': int(time.time()), 'tables': {}}
            with connect() as source:
                for table in tables:
                    snapshot['tables'][table] = [dict(row) for row in source.execute(f'SELECT * FROM {table}')]
            content = json.dumps(snapshot, ensure_ascii=False, indent=2).encode('utf-8')
            return send_file(io.BytesIO(content), as_attachment=True,
                download_name='backup-datafy.json', mimetype='application/json')
        # SQLite backup API includes committed WAL transactions.
        dest = sqlite3.connect(':memory:')
        try:
            with connect() as source:
                source.raw.backup(dest)
            content = dest.serialize()
        finally:
            dest.close()
        return send_file(io.BytesIO(content), as_attachment=True, download_name='backup-datafy.sqlite3', mimetype='application/octet-stream')

    return panel, webhook

def create_render_app(db_path=None, settings=None):
    panel, webhook = create_apps(db_path=db_path, settings=settings)
    panel.add_url_rule('/webhook/datafy', view_func=webhook.view_functions['receive'], methods=['POST'])
    return panel

# Gunicorn entry point used by Render.
app = create_render_app() if os.environ.get('RENDER') else None

if __name__ == '__main__':
    settings_path = ROOT / 'config.json'
    settings = json.loads(settings_path.read_text(encoding='utf-8-sig')) if settings_path.exists() else {}
    panel, webhook = create_apps(settings=settings)
    pp, wp = int(settings.get('panel_port',8765)), int(settings.get('webhook_port',8766))
    threading.Thread(target=lambda: serve(webhook,host='127.0.0.1',port=wp,threads=4),daemon=True).start()
    print(f'Painel: http://127.0.0.1:{pp} | Receptor local: http://127.0.0.1:{wp}/webhook/datafy', flush=True)
    print('Para encerrar, pressione Ctrl+C. Não publique a porta do painel.', flush=True)
    serve(panel,host='127.0.0.1',port=pp,threads=4)
